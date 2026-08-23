import json
import os
import re
from pathlib import Path

from openai import OpenAI
from config import DEEPSEEK_API_KEY, DEEPSEEK_BASE_URL, DEEPSEEK_MODEL

client = OpenAI(api_key=DEEPSEEK_API_KEY, base_url=DEEPSEEK_BASE_URL)


def extract_text(path: Path, ext: str) -> str:
    ext = ext.lower().lstrip(".")
    if ext in ("md", "txt"):
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == "pdf":
        return _extract_pdf(path)
    if ext == "docx":
        return _extract_docx(path)
    if ext == "doc":
        return _extract_legacy_doc(path)
    raise ValueError(f"Unsupported file type: {ext}")


def _extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        parts = []
        for page in reader.pages:
            t = page.extract_text()
            if t and t.strip():
                parts.append(t.strip())
        return "\n\n".join(parts)
    except ImportError:
        raise RuntimeError("请安装 pypdf：pip install pypdf")
    except Exception as e:
        raise RuntimeError(f"PDF 解析失败：{e}")


def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
        doc = Document(str(path))
        parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if row_text:
                    parts.append(row_text)
        return "\n".join(parts)
    except ImportError:
        raise RuntimeError("请安装 python-docx：pip install python-docx")
    except Exception as e:
        raise RuntimeError(f"DOCX 解析失败：{e}")


def _extract_legacy_doc(path: Path) -> str:
    """Extract Word 97-2003 .doc text through an installed Microsoft Word."""
    if os.name != "nt":
        raise RuntimeError("旧版 DOC 仅支持在安装了 Microsoft Word 的 Windows 环境解析；请另存为 DOCX 后上传")

    try:
        import pythoncom
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("解析旧版 DOC 需要 pywin32 和 Microsoft Word；请另存为 DOCX 后上传") from exc

    word = None
    document = None
    pythoncom.CoInitialize()
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        word.DisplayAlerts = 0
        # Disable AutoOpen and other document macros for imported files.
        word.AutomationSecurity = 3
        document = word.Documents.Open(
            str(path.resolve()),
            ConfirmConversions=False,
            ReadOnly=True,
            AddToRecentFiles=False,
        )
        text = str(document.Content.Text or "")
        # Word uses carriage returns for paragraphs and \x07 for table cells.
        return text.replace("\r", "\n").replace("\x07", "\t").strip()
    except Exception as exc:
        raise RuntimeError(f"DOC 解析失败：{exc}；可尝试在 Word 中另存为 DOCX 后上传") from exc
    finally:
        if document is not None:
            try:
                document.Close(False)
            except Exception:
                pass
        if word is not None:
            try:
                word.Quit()
            except Exception:
                pass
        pythoncom.CoUninitialize()


def generate_note_metadata(text: str, title: str) -> dict:
    """Use LLM to produce summary, category, tags. Returns dict."""
    try:
        resp = client.chat.completions.create(
            model=DEEPSEEK_MODEL,
            max_tokens=500,
            messages=[
                {
                    "role": "system",
                    "content": "你是笔记整理专家，只输出JSON，不要解释。",
                },
                {
                    "role": "user",
                    "content": (
                        f"请分析笔记「{title}」，输出JSON：\n"
                        '{"summary": "2-3句话核心摘要", '
                        '"category": "技术/商业/科学/人文/健康/其他 选一个", '
                        '"tags": ["标签1", "标签2", "标签3"]}\n\n'
                        f"{text[:3000]}"
                    ),
                },
            ],
        )
        raw = resp.choices[0].message.content or "{}"
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if not match:
            return {"summary": "", "category": "其他", "tags": []}
        data = json.loads(match.group())
        return {
            "summary": data.get("summary", ""),
            "category": data.get("category", "其他"),
            "tags": [str(t) for t in data.get("tags", [])[:6]],
        }
    except Exception:
        return {"summary": "", "category": "其他", "tags": []}
